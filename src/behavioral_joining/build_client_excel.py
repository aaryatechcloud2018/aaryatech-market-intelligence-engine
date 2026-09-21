"""build_client_excel.py -- ABIDS_Client_Intelligence_Dashboard.xlsx"""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
GOLD = "9C6B1F"
WHITE = "FFFFFF"
HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
HEADER_FONT = Font(color=WHITE, bold=True)
TITLE_FONT = Font(size=16, bold=True, color=NAVY)
KPI_FONT = Font(size=22, bold=True, color=NAVY)
KPI_LABEL_FONT = Font(size=11, color="555555")
CARD_FILL = PatternFill(start_color="EFF3FA", end_color="EFF3FA", fill_type="solid")
CARD_BORDER = Border(*[Side(style="thin", color="C6D2EA")] * 4)


def _write_df(wb, name, df, heatmap_cols=None):
    ws = wb.create_sheet(name)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    for col in ws.columns:
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 50)
    if heatmap_cols:
        n_rows = len(df) + 1
        for col_name in heatmap_cols:
            if col_name not in df.columns:
                continue
            col_idx = df.columns.get_loc(col_name) + 1
            col_letter = get_column_letter(col_idx)
            rng = f"{col_letter}2:{col_letter}{n_rows}"
            rule = ColorScaleRule(
                start_type="min", start_color="63BE7B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="F8696B",
            )
            ws.conditional_formatting.add(rng, rule)
    return ws


def _kpi_card(ws, row, col, label, value):
    ws.cell(row=row, column=col, value=value).font = KPI_FONT
    ws.cell(row=row + 1, column=col, value=label).font = KPI_LABEL_FONT
    for r in (row, row + 1):
        ws.cell(row=r, column=col).fill = CARD_FILL
        ws.cell(row=r, column=col).border = CARD_BORDER
        ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")


def build_workbook(out_path, tables, bpi_stage, bpi_detail, apps_summary):
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("00_EXECUTIVE_DASHBOARD")
    ws["A1"] = "ABIDS Client Intelligence Dashboard"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Demonstration dataset -- not real client historical data."
    ws["A2"].font = Font(italic=True, color="A6231A")
    for i, (label, value) in enumerate(apps_summary.items()):
        _kpi_card(ws, 4, 2 + i * 3, label, value)
    ws.column_dimensions["A"].width = 4
    for c in range(2, 2 + len(apps_summary) * 3):
        ws.column_dimensions[get_column_letter(c)].width = 16

    ws["A8"] = "Top Priority Stages (BPI MVP)"
    ws["A8"].font = Font(bold=True, size=13, color=GOLD)
    r0 = 9
    ws.cell(row=r0, column=1, value="Stage")
    ws.cell(row=r0, column=2, value="BPI MVP Score")
    ws.cell(row=r0, column=3, value="Band")
    for c in (1, 2, 3):
        ws.cell(row=r0, column=c).font = HEADER_FONT
        ws.cell(row=r0, column=c).fill = HEADER_FILL
    for i, row in enumerate(bpi_stage.head(5).itertuples(), start=1):
        ws.cell(row=r0 + i, column=1, value=row.journey_stage)
        ws.cell(row=r0 + i, column=2, value=row.bpi_mvp_score)
        ws.cell(row=r0 + i, column=3, value=row.relative_priority_band)

    _write_df(wb, "01_stage_summary", tables["stage_summary"], heatmap_cols=["dropoff_rate", "behavior_penetration_rate"])
    _write_df(wb, "02_stage_behavior_matrix", tables["stage_behavior_matrix"], heatmap_cols=["behavior_share_within_stage", "dropoff_rate_associated"])
    _write_df(wb, "03_stage_scenario_matrix", tables["stage_scenario_matrix"], heatmap_cols=["scenario_share_within_stage", "associated_dropoff_rate"])
    _write_df(wb, "04_stage_behavior_scenario_map", tables["stage_behavior_scenario_map"])
    _write_df(wb, "05_uncertainty_stage_summary", tables["uncertainty_stage_summary"], heatmap_cols=["uncertainty_score"])
    _write_df(wb, "06_hypothesis_visual_summary", tables["hypothesis_visual_summary"], heatmap_cols=["absolute_difference"])
    _write_df(wb, "07_stage_priority_bpi_mvp", bpi_stage, heatmap_cols=["bpi_mvp_score"])
    _write_df(wb, "08_scenario_behavior_priority", bpi_detail, heatmap_cols=["bpi_mvp_score"])

    wb.save(out_path)
    return out_path
